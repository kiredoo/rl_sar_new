#!/usr/bin/env python3
"""Aggregate per-config metrics into one wide table (csv + xlsx).

Reads configs.tsv at the experiments root, looks up each config's
runs/<run_id>/derived/summary.tsv, and emits:

  derived_cross/aggregate.csv
  derived_cross/aggregate.xlsx   (if openpyxl available)

Each row = one logical config (baseline, stockpile, hybrid-unfixed, ...).
Missing configs (run_id == "-") appear as "(not yet run)" so the table
shape stays stable as new runs land.

Baseline column: the first config whose run_id is populated is treated as
the baseline; %-delta columns (vs baseline) are added for minflt, rss_kb
growth, and rss_kb steady-state (pre).

Usage:
    tools/aggregate_configs.py               # default paths
    tools/aggregate_configs.py --root <dir>  # alt root
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

METRIC_COLS = [
    "minflt_delta", "majflt_delta", "rss_kb_delta",
    "vol_ctx_delta", "invol_ctx_delta", "rss_kb_pre", "nprocs", "aborts",
]

# Single source of truth for what each column means. Rendered into the
# xlsx "metrics" sheet and derived_cross/definitions.md.
# Fields: column name, unit, source (file/field), definition, sign convention.
METRIC_DEFS: list[dict[str, str]] = [
    {
        "column": "name",
        "unit": "—",
        "source": "configs.tsv",
        "definition": "Logical config label. Stable across runs (e.g. 'baseline', 'stockpile', 'hybrid-unfixed', 'hybrid-fixed-v1', ...).",
        "sign": "—",
    },
    {
        "column": "run_id",
        "unit": "—",
        "source": "configs.tsv → runs/<run_id>/",
        "definition": "Immutable run directory this config's numbers come from. Format: YYYYMMDD-HHMM_<slug> (going-forward) or <ISO-date>T<HH-MM>_<slug> (pre-2026-04-29 grandfathered). See docs/CONVENTIONS.md.",
        "sign": "—",
    },
    {
        "column": "arm",
        "unit": "—",
        "source": "configs.tsv → runs/<run_id>/raw/<arm>/",
        "definition": "Arm label inside the run (A, B, C, D...). One run can hold multiple arms of the same experiment.",
        "sign": "—",
    },
    {
        "column": "minflt_delta",
        "unit": "page faults (count)",
        "source": "/proc/<pid>/stat field 10; post − pre, summed over all matched pids",
        "definition": ("Minor page faults during the bag-play phase. A minor fault = page not resident but no disk I/O "
                       "(demand-zero, COW break, anon-page touch). Heavy allocation + touch produces lots of these even "
                       "when RSS is flat, because freed-and-reallocated pages re-fault."),
        "sign": "lower is better (less allocator churn)",
    },
    {
        "column": "majflt_delta",
        "unit": "page faults (count)",
        "source": "/proc/<pid>/stat field 12; post − pre, summed over all matched pids",
        "definition": ("Major page faults during the bag-play phase. A major fault = page required disk I/O "
                       "(swap-in, file-backed mmap miss). For our fully-RAM workload this should stay ~0; non-zero "
                       "indicates memory pressure or swap."),
        "sign": "lower is better; ≈0 expected on a healthy box",
    },
    {
        "column": "rss_kb_delta",
        "unit": "kB (kibibytes, per Linux /proc convention)",
        "source": "/proc/<pid>/status VmRSS; post − pre, summed over all matched pids",
        "definition": ("Growth of resident set size across all Autoware processes during the bag-play phase. "
                       "Captures allocator behaviour under load: an allocator that hoards freed pages will show "
                       "higher growth than one that returns memory to the OS."),
        "sign": "lower is better (less runtime growth). Negative = RSS shrank (unusual under load).",
    },
    {
        "column": "vol_ctx_delta",
        "unit": "context switches (count)",
        "source": "/proc/<pid>/status voluntary_ctxt_switches; post − pre, summed over all matched pids",
        "definition": ("Voluntary context switches = thread blocked itself (I/O wait, condvar, sleep). Large "
                       "differences between allocator arms can indicate lock contention inside the allocator "
                       "(e.g. shared TLSF fallback lock)."),
        "sign": "informational; interpret in context of allocator's locking model",
    },
    {
        "column": "invol_ctx_delta",
        "unit": "context switches (count)",
        "source": "/proc/<pid>/status nonvoluntary_ctxt_switches; post − pre, summed over all matched pids",
        "definition": ("Involuntary context switches = thread preempted by scheduler (quantum expired, higher-prio "
                       "task). Indirectly reflects CPU pressure during the phase."),
        "sign": "informational",
    },
    {
        "column": "rss_kb_pre",
        "unit": "kB",
        "source": "/proc/<pid>/status VmRSS at pre-snapshot; summed over all matched pids",
        "definition": ("Steady-state resident memory after Autoware bringup (~75 s launch + idle detection), "
                       "BEFORE bag play starts. Measures the allocator's post-init footprint — mapping tables, "
                       "model weights, preallocated pools."),
        "sign": "lower is better (smaller steady-state footprint)",
    },
    {
        "column": "nprocs",
        "unit": "process count",
        "source": "lines of snapshot_rss.sh output (one per matched pid)",
        "definition": ("Number of Autoware-related processes the pgrep filter captured at pre-snapshot time. "
                       "Should match across A/B arms for an apples-to-apples comparison. Current filter: "
                       "ros-args|rviz2|ros2 launch autoware|robot_state_publisher|web_server.py|logging_simulator."),
        "sign": "must match between arms; mismatch invalidates the comparison",
    },
    {
        "column": "aborts",
        "unit": "count",
        "source": "zgrep/grep 'exit code -6' over raw/<arm>/aw_<arm>_*.log[.gz]",
        "definition": ("Number of component_container processes that SIGABRT'd during bag play. Any non-zero "
                       "value means this arm's delta metrics include dying-process allocations and are NOT "
                       "directly comparable to a clean arm's delta — steady-state rss_kb_pre remains valid."),
        "sign": "must be 0 for arm-vs-arm delta comparison; empty = pre-aborts-tracking run (backfill via compute_summary.sh)",
    },
    {
        "column": "minflt_delta_vs_base_%",
        "unit": "percent",
        "source": "derived: (this.minflt_delta − baseline.minflt_delta) / baseline.minflt_delta",
        "definition": "Percent change in bag-phase minor faults vs the baseline config (first non-empty row in configs.tsv).",
        "sign": "negative = better than baseline",
    },
    {
        "column": "rss_kb_delta_vs_base_%",
        "unit": "percent",
        "source": "derived: (this.rss_kb_delta − baseline.rss_kb_delta) / baseline.rss_kb_delta",
        "definition": "Percent change in bag-phase RSS growth vs baseline.",
        "sign": "negative = better than baseline",
    },
    {
        "column": "rss_kb_pre_vs_base_%",
        "unit": "percent",
        "source": "derived: (this.rss_kb_pre − baseline.rss_kb_pre) / baseline.rss_kb_pre",
        "definition": "Percent change in steady-state RSS vs baseline.",
        "sign": "negative = smaller footprint than baseline",
    },
    {
        "column": "notes",
        "unit": "—",
        "source": "configs.tsv",
        "definition": "Free-form label from the registry (allocator variant, branch, any relevant flags).",
        "sign": "—",
    },
]

PHASE_NOTE = (
    "Phase definitions:\n"
    "  pre-snapshot  — captured AFTER bringup settles (~75 s wait + idle detection), BEFORE `ros2 bag play` starts.\n"
    "  post-snapshot — captured AFTER `ros2 bag play` exits, BEFORE SIGINT to the launch.\n"
    "  delta         — post − pre, summed per metric across all matched pids (the bag-phase workload only).\n"
    "  All counters are per-process kernel-maintained; we only diff, never sample rates.\n"
)


def read_configs(path: Path) -> list[dict]:
    rows = []
    with path.open() as f:
        reader = csv.DictReader(
            (ln for ln in f if ln.strip() and not ln.lstrip().startswith("#")),
            delimiter="\t",
        )
        for r in reader:
            rows.append({k.strip(): (v or "").strip() for k, v in r.items()})
    return rows


SUMMARY_TSV_SCHEMA = 1


def _read_lines_skip_comments(path: Path) -> tuple[int | None, list[str]]:
    """Return (schema_version, non_comment_lines).

    schema_version is parsed from the first `# schema=N` line if present,
    else None (grandfathered pre-schema runs). Other `#`-prefixed lines
    are treated as free-form comments and discarded.
    """
    schema: int | None = None
    out: list[str] = []
    with path.open() as f:
        for ln in f:
            s = ln.rstrip("\n")
            if s.startswith("#"):
                m = s.lstrip("#").strip()
                if m.startswith("schema="):
                    try:
                        schema = int(m.split("=", 1)[1])
                    except ValueError:
                        # malformed schema line — leave schema as None
                        # so the caller can decide whether to abort
                        pass
                continue
            out.append(s)
    return schema, out


def read_summary(run_dir: Path, arm: str) -> dict | None:
    tsv = run_dir / "derived" / "summary.tsv"
    if not tsv.exists():
        return None
    schema, lines = _read_lines_skip_comments(tsv)
    if schema is not None and schema != SUMMARY_TSV_SCHEMA:
        raise SystemExit(
            f"summary.tsv schema mismatch in {tsv}: file=schema={schema}, "
            f"tool expects schema={SUMMARY_TSV_SCHEMA}. "
            "Bump the consumer or regenerate the file."
        )
    if not lines:
        return None
    header = lines[0].split("\t")
    for ln in lines[1:]:
        parts = ln.split("\t")
        if parts and parts[0] == arm:
            return dict(zip(header, parts))
    return None


def pct(v: str | float | None, base: str | float | None) -> str:
    try:
        v_f = float(v); b_f = float(base)
    except (TypeError, ValueError):
        return ""
    if b_f == 0:
        return ""
    return f"{(v_f - b_f) * 100.0 / b_f:+.1f}%"


def build_table(root: Path) -> tuple[list[str], list[list[str]]]:
    configs = read_configs(root / "configs.tsv")

    records: list[dict] = []
    for c in configs:
        rec = {"name": c["name"], "run_id": c["run_id"],
               "arm": c["arm"], "notes": c["notes"]}
        if c["run_id"] in ("", "-"):
            rec["_missing"] = True
        else:
            s = read_summary(root / "runs" / c["run_id"], c["arm"])
            if s is None:
                rec["_missing"] = True
                rec["notes"] = (rec["notes"] + " [summary.tsv missing]").strip()
            else:
                for col in METRIC_COLS:
                    rec[col] = s.get(col, "")
        records.append(rec)

    baseline = next((r for r in records if not r.get("_missing")), None)

    header = ["name", "run_id", "arm"] + METRIC_COLS + [
        "minflt_delta_vs_base_%", "rss_kb_delta_vs_base_%",
        "rss_kb_pre_vs_base_%", "notes",
    ]
    rows: list[list[str]] = []
    for r in records:
        if r.get("_missing"):
            rows.append([r["name"], "(not yet run)", ""] + [""] * len(METRIC_COLS)
                        + ["", "", "", r["notes"]])
            continue
        pct_mi = pct(r.get("minflt_delta"), baseline and baseline.get("minflt_delta"))
        pct_rd = pct(r.get("rss_kb_delta"), baseline and baseline.get("rss_kb_delta"))
        pct_rp = pct(r.get("rss_kb_pre"), baseline and baseline.get("rss_kb_pre"))
        if baseline is r:
            pct_mi = pct_rd = pct_rp = "(baseline)"
        rows.append(
            [r["name"], r["run_id"], r["arm"]]
            + [r.get(c, "") for c in METRIC_COLS]
            + [pct_mi, pct_rd, pct_rp, r["notes"]]
        )
    return header, rows


def write_csv(out: Path, header: list[str], rows: list[list[str]]) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def write_definitions_md(out: Path) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    lines: list[str] = []
    lines.append("# aggregate.{csv,xlsx} — column definitions\n")
    lines.append("Generated by `tools/aggregate_configs.py`. Do not hand-edit — "
                 "update `METRIC_DEFS` in the script.\n")
    lines.append("## Phase definitions\n")
    lines.append("```")
    lines.append(PHASE_NOTE.rstrip())
    lines.append("```\n")
    lines.append("## Columns\n")
    lines.append("| column | unit | source | definition | sign convention |")
    lines.append("|---|---|---|---|---|")
    for d in METRIC_DEFS:
        lines.append("| `{column}` | {unit} | {source} | {definition} | {sign} |".format(**d))
    lines.append("")
    out.write_text("\n".join(lines))


def write_xlsx(out: Path, header: list[str], rows: list[list[str]]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "aggregate"
    ws.append(header)
    # Definitions sheet so the xlsx is self-contained.
    ws_def = wb.create_sheet("metrics")
    ws_def.append(["column", "unit", "source", "definition", "sign convention"])
    for d in METRIC_DEFS:
        ws_def.append([d["column"], d["unit"], d["source"], d["definition"], d["sign"]])
    for col_idx, width in enumerate([24, 22, 44, 90, 44], 1):
        ws_def.column_dimensions[ws_def.cell(row=1, column=col_idx).column_letter].width = width
    for cell in ws_def[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(horizontal="center")
    for row in ws_def.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws_def.freeze_panes = "A2"
    # Phase note on a third sheet.
    ws_phase = wb.create_sheet("phases")
    for ln in PHASE_NOTE.rstrip().splitlines():
        ws_phase.append([ln])
    ws_phase.column_dimensions["A"].width = 120
    hdr_font = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="DDDDDD")
    for cell in ws[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        # coerce numerics so Excel treats them as numbers
        coerced = []
        for v in r:
            if v in ("", "(not yet run)", "(baseline)") or str(v).endswith("%"):
                coerced.append(v)
            else:
                try:
                    coerced.append(int(v))
                except (TypeError, ValueError):
                    try:
                        coerced.append(float(v))
                    except (TypeError, ValueError):
                        coerced.append(v)
        ws.append(coerced)
    for col_idx, col_name in enumerate(header, 1):
        max_len = max([len(str(col_name))]
                      + [len(str(row[col_idx - 1])) for row in rows])
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 40)
    ws.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return True


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--root", type=Path,
                   default=Path(__file__).resolve().parent.parent,
                   help="experiments root (default: parent of tools/)")
    p.add_argument("--out-dir", type=Path, default=None,
                   help="output dir (default: <root>/derived_cross/)")
    args = p.parse_args()

    header, rows = build_table(args.root)
    out_dir = args.out_dir or (args.root / "derived_cross")
    csv_path = out_dir / "aggregate.csv"
    xlsx_path = out_dir / "aggregate.xlsx"
    defs_path = out_dir / "definitions.md"
    write_csv(csv_path, header, rows)
    xlsx_ok = write_xlsx(xlsx_path, header, rows)
    write_definitions_md(defs_path)

    print(f"wrote {csv_path}")
    print(f"wrote {xlsx_path}" if xlsx_ok else "xlsx skipped (openpyxl not installed)")
    print(f"wrote {defs_path}")
    print()
    widths = [max(len(str(h)), max((len(str(r[i])) for r in rows), default=0))
              for i, h in enumerate(header)]
    fmt = "  ".join(f"{{:<{w}}}" for w in widths)
    print(fmt.format(*header))
    print(fmt.format(*["-" * w for w in widths]))
    for r in rows:
        print(fmt.format(*[str(x) for x in r]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
